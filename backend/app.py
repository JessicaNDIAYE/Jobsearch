import os
import uuid
import json
import hashlib
import sqlite3
import urllib.request
import urllib.parse
from datetime import datetime
from io import BytesIO
from contextlib import contextmanager

from flask import Flask, request, jsonify, send_file, g
from flask_cors import CORS
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)

allowed_origins = ["http://localhost:5173", "http://localhost:3000"]
frontend_url = os.environ.get("FRONTEND_URL", "")
if frontend_url:
    allowed_origins.append(frontend_url)

CORS(app, origins=allowed_origins if frontend_url else "*",
     methods=["GET", "POST", "PUT", "DELETE", "OPTIONS"],
     allow_headers=["Content-Type", "Authorization"])

# ---------------------------------------------------------------------------
# DB — PostgreSQL (prod) or SQLite (local)
# ---------------------------------------------------------------------------

DATABASE_URL = os.environ.get("DATABASE_URL", "")
USE_PG = bool(DATABASE_URL)

if USE_PG:
    import psycopg2
    import psycopg2.extras
else:
    DB_PATH = os.environ.get("DB_PATH", os.path.join(os.path.dirname(__file__), "jobtracker.db"))


def get_db():
    if "db" not in g:
        if USE_PG:
            conn = psycopg2.connect(DATABASE_URL, cursor_factory=psycopg2.extras.RealDictCursor)
            g.db = conn
        else:
            conn = sqlite3.connect(DB_PATH)
            conn.row_factory = sqlite3.Row
            g.db = conn
    return g.db


@app.teardown_appcontext
def close_db(exc=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def query(sql, params=(), fetchone=False, fetchall=False, commit=False):
    """Unified query helper for both PG and SQLite."""
    db = get_db()
    if USE_PG:
        # Translate ? placeholders to %s for psycopg2
        pg_sql = sql.replace("?", "%s")
        cur = db.cursor()
        cur.execute(pg_sql, params)
        result = None
        if fetchone:
            result = cur.fetchone()
        elif fetchall:
            result = cur.fetchall()
        if commit:
            db.commit()
        cur.close()
        return result
    else:
        cur = db.execute(sql, params)
        result = None
        if fetchone:
            result = cur.fetchone()
        elif fetchall:
            result = cur.fetchall()
        if commit:
            db.commit()
        return result


def init_db():
    if USE_PG:
        conn = psycopg2.connect(DATABASE_URL)
        cur = conn.cursor()
        cur.execute("""
            CREATE TABLE IF NOT EXISTS users (
                id TEXT PRIMARY KEY,
                username TEXT NOT NULL,
                email TEXT UNIQUE NOT NULL,
                password_hash TEXT NOT NULL,
                created_at TEXT NOT NULL
            )
        """)
        cur.execute("""
            CREATE TABLE IF NOT EXISTS sessions (
                token TEXT PRIMARY KEY,
                user_id TEXT NOT NULL,
                created_at TEXT NOT NULL
            )
        """)
        cur.execute("""
            CREATE TABLE IF NOT EXISTS saved_jobs (
                id TEXT PRIMARY KEY,
                user_id TEXT NOT NULL,
                company TEXT,
                type TEXT,
                industry TEXT,
                location TEXT,
                notes TEXT,
                job_title TEXT,
                job_url TEXT,
                salary TEXT,
                status TEXT DEFAULT 'saved',
                tags TEXT DEFAULT '[]',
                created_at TEXT NOT NULL
            )
        """)
        conn.commit()
        cur.close()
        conn.close()
    else:
        conn = sqlite3.connect(DB_PATH)
        conn.executescript("""
            CREATE TABLE IF NOT EXISTS users (
                id TEXT PRIMARY KEY,
                username TEXT NOT NULL,
                email TEXT UNIQUE NOT NULL,
                password_hash TEXT NOT NULL,
                created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS sessions (
                token TEXT PRIMARY KEY,
                user_id TEXT NOT NULL,
                created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS saved_jobs (
                id TEXT PRIMARY KEY,
                user_id TEXT NOT NULL,
                company TEXT,
                type TEXT,
                industry TEXT,
                location TEXT,
                notes TEXT,
                job_title TEXT,
                job_url TEXT,
                salary TEXT,
                status TEXT DEFAULT 'saved',
                tags TEXT DEFAULT '[]',
                created_at TEXT NOT NULL
            );
        """)
        conn.commit()
        conn.close()


# ---------------------------------------------------------------------------
# Demo seed data
# ---------------------------------------------------------------------------

DEMO_JOBS = [
    {"company": "Agicap",     "job_title": "Data Analyst Junior",  "type": "Scale-up",        "industry": "Fintech/SaaS",    "location": "Lyon",               "status": "saved", "salary": "", "job_url": "", "notes": "Scale-up fintech en forte croissance, bonne culture data", "tags": []},
    {"company": "Capgemini",  "job_title": "Junior AI Engineer",   "type": "Grande entreprise","industry": "Consulting/AI",   "location": "Lyon",               "status": "saved", "salary": "", "job_url": "", "notes": "Missions variées en AI/Data, bonne école pour débuter",   "tags": []},
    {"company": "DataGenius", "job_title": "Data Scientist",       "type": "PME",             "industry": "Data Science",    "location": "Lyon (Villeurbanne)","status": "saved", "salary": "", "job_url": "", "notes": "Petite structure avec projets ML concrets",               "tags": []},
    {"company": "MILA",       "job_title": "Research Assistant",   "type": "Research Institute","industry": "AI Research",   "location": "Montreal",           "status": "saved", "salary": "", "job_url": "", "notes": "Institut de recherche en IA de renommée mondiale",        "tags": []},
    {"company": "Shopify",    "job_title": "Junior Data Scientist","type": "Big Tech",        "industry": "Data & e-commerce","location": "Canada (remote)",   "status": "saved", "salary": "", "job_url": "", "notes": "Télétravail complet, stack data moderne",                 "tags": []},
]

# ---------------------------------------------------------------------------
# Auth helpers
# ---------------------------------------------------------------------------

def hash_password(p): return hashlib.sha256(p.encode()).hexdigest()


def get_current_user():
    auth = request.headers.get("Authorization", "")
    if not auth.startswith("Bearer "):
        return None
    token = auth[7:]
    row = query(
        "SELECT u.* FROM sessions s JOIN users u ON s.user_id = u.id WHERE s.token = ?",
        (token,), fetchone=True
    )
    return dict(row) if row else None


def require_auth(f):
    from functools import wraps
    @wraps(f)
    def wrapper(*args, **kwargs):
        user = get_current_user()
        if not user:
            return jsonify({"error": "Unauthorized"}), 401
        return f(user, *args, **kwargs)
    return wrapper

# ---------------------------------------------------------------------------
# Routes — Auth
# ---------------------------------------------------------------------------

@app.route("/api/register", methods=["POST"])
def register():
    data = request.get_json()
    username = (data.get("username") or "").strip()
    email    = (data.get("email") or "").strip().lower()
    password =  data.get("password") or ""
    if not username or not email or not password:
        return jsonify({"error": "All fields are required"}), 400

    existing = query("SELECT id FROM users WHERE email = ?", (email,), fetchone=True)
    if existing:
        return jsonify({"error": "Email already registered"}), 409

    user_id = str(uuid.uuid4())
    now     = datetime.utcnow().isoformat()
    query("INSERT INTO users (id, username, email, password_hash, created_at) VALUES (?,?,?,?,?)",
          (user_id, username, email, hash_password(password), now), commit=True)

    for job in DEMO_JOBS:
        query("""INSERT INTO saved_jobs
                 (id,user_id,company,type,industry,location,notes,job_title,job_url,salary,status,tags,created_at)
                 VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
              (str(uuid.uuid4()), user_id, job["company"], job["type"], job["industry"],
               job["location"], job["notes"], job["job_title"], job["job_url"],
               job["salary"], job["status"], json.dumps(job["tags"]), now), commit=True)

    token = str(uuid.uuid4())
    query("INSERT INTO sessions (token, user_id, created_at) VALUES (?,?,?)",
          (token, user_id, now), commit=True)
    return jsonify({"token": token, "username": username, "email": email}), 201


@app.route("/api/login", methods=["POST"])
def login():
    data  = request.get_json()
    email = (data.get("email") or "").strip().lower()
    pw    =  data.get("password") or ""
    user  = query("SELECT * FROM users WHERE email = ? AND password_hash = ?",
                  (email, hash_password(pw)), fetchone=True)
    if not user:
        return jsonify({"error": "Invalid email or password"}), 401
    token = str(uuid.uuid4())
    now   = datetime.utcnow().isoformat()
    query("INSERT INTO sessions (token, user_id, created_at) VALUES (?,?,?)",
          (token, user["id"], now), commit=True)
    return jsonify({"token": token, "username": user["username"], "email": user["email"]})


@app.route("/api/logout", methods=["POST"])
@require_auth
def logout(user):
    token = request.headers.get("Authorization", "")[7:]
    query("DELETE FROM sessions WHERE token = ?", (token,), commit=True)
    return jsonify({"ok": True})

# ---------------------------------------------------------------------------
# Routes — Jobs
# ---------------------------------------------------------------------------

def row_to_job(row):
    d = dict(row)
    try:    d["tags"] = json.loads(d.get("tags") or "[]")
    except: d["tags"] = []
    return d


@app.route("/api/jobs", methods=["GET"])
@require_auth
def get_jobs(user):
    rows = query("SELECT * FROM saved_jobs WHERE user_id = ? ORDER BY created_at DESC",
                 (user["id"],), fetchall=True)
    return jsonify([row_to_job(r) for r in rows])


@app.route("/api/jobs", methods=["POST"])
@require_auth
def create_job(user):
    data   = request.get_json()
    job_id = str(uuid.uuid4())
    now    = datetime.utcnow().isoformat()
    query("""INSERT INTO saved_jobs
             (id,user_id,company,type,industry,location,notes,job_title,job_url,salary,status,tags,created_at)
             VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
          (job_id, user["id"], data.get("company",""), data.get("type",""),
           data.get("industry",""), data.get("location",""), data.get("notes",""),
           data.get("job_title",""), data.get("job_url",""), data.get("salary",""),
           data.get("status","saved"), json.dumps(data.get("tags") or []), now), commit=True)
    row = query("SELECT * FROM saved_jobs WHERE id = ?", (job_id,), fetchone=True)
    return jsonify(row_to_job(row)), 201


@app.route("/api/jobs/<job_id>", methods=["PUT"])
@require_auth
def update_job(user, job_id):
    existing = query("SELECT id FROM saved_jobs WHERE id = ? AND user_id = ?",
                     (job_id, user["id"]), fetchone=True)
    if not existing:
        return jsonify({"error": "Not found"}), 404
    data = request.get_json()
    query("""UPDATE saved_jobs SET
             company=?,type=?,industry=?,location=?,notes=?,
             job_title=?,job_url=?,salary=?,status=?,tags=?
             WHERE id=? AND user_id=?""",
          (data.get("company",""), data.get("type",""), data.get("industry",""),
           data.get("location",""), data.get("notes",""), data.get("job_title",""),
           data.get("job_url",""), data.get("salary",""), data.get("status","saved"),
           json.dumps(data.get("tags") or []), job_id, user["id"]), commit=True)
    row = query("SELECT * FROM saved_jobs WHERE id = ?", (job_id,), fetchone=True)
    return jsonify(row_to_job(row))


@app.route("/api/jobs/<job_id>", methods=["DELETE"])
@require_auth
def delete_job(user, job_id):
    existing = query("SELECT id FROM saved_jobs WHERE id = ? AND user_id = ?",
                     (job_id, user["id"]), fetchone=True)
    if not existing:
        return jsonify({"error": "Not found"}), 404
    query("DELETE FROM saved_jobs WHERE id = ?", (job_id,), commit=True)
    return jsonify({"ok": True})

# ---------------------------------------------------------------------------
# Routes — Stats
# ---------------------------------------------------------------------------

@app.route("/api/stats", methods=["GET"])
@require_auth
def get_stats(user):
    rows = query("SELECT status, location FROM saved_jobs WHERE user_id = ?",
                 (user["id"],), fetchall=True)
    by_status, by_location = {}, {}
    for r in rows:
        s = r["status"] or "saved"
        by_status[s] = by_status.get(s, 0) + 1
        loc = r["location"] or "Unknown"
        by_location[loc] = by_location.get(loc, 0) + 1
    return jsonify({"total": len(rows), "by_status": by_status, "by_location": by_location})

# ---------------------------------------------------------------------------
# Routes — Excel Export
# ---------------------------------------------------------------------------

STATUS_COLORS = {"saved":"3B82F6","applied":"F97316","interview":"22C55E","offer":"A855F7","rejected":"EF4444"}
DARK_BG="1a1a2e"; HEADER_BG="e94560"; ROW_ALT="16213e"; ROW_MAIN="0f3460"
TEXT_WHITE="FFFFFF"; TEXT_LIGHT="e2e8f0"

def make_border():
    s = Side(style="thin", color="2d3748")
    return Border(left=s, right=s, top=s, bottom=s)


@app.route("/api/jobs/export", methods=["GET"])
@require_auth
def export_jobs(user):
    rows = query("SELECT * FROM saved_jobs WHERE user_id = ? ORDER BY created_at DESC",
                 (user["id"],), fetchall=True)
    wb = Workbook(); ws = wb.active; ws.title = "Job Tracker"

    ws.merge_cells("A1:J1")
    c = ws["A1"]
    c.value = f"JobTracker — {user['username']} — {datetime.utcnow().strftime('%Y-%m-%d')}"
    c.font = Font(name="Calibri", bold=True, size=14, color=TEXT_WHITE)
    c.fill = PatternFill("solid", fgColor=HEADER_BG)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    headers = ["Company","Job Title","Type","Industry","Location","Status","Salary","Notes","URL","Saved On"]
    widths  = [22,26,18,20,18,12,14,35,35,16]
    for ci, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=2, column=ci, value=h)
        cell.font = Font(name="Calibri", bold=True, size=11, color=TEXT_WHITE)
        cell.fill = PatternFill("solid", fgColor="2d1b4e")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = make_border()
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.row_dimensions[2].height = 22

    for rn, job in enumerate(rows, 3):
        bg = ROW_ALT if rn % 2 == 0 else ROW_MAIN
        sc = STATUS_COLORS.get((job["status"] or "saved").lower(), "3B82F6")
        vals = [job["company"],job["job_title"],job["type"],job["industry"],job["location"],
                (job["status"] or "saved").capitalize(),job["salary"],job["notes"],
                job["job_url"],(job["created_at"] or "")[:10]]
        for ci, v in enumerate(vals, 1):
            cell = ws.cell(row=rn, column=ci, value=v)
            cell.font = Font(name="Calibri", size=10, color=TEXT_LIGHT)
            cell.alignment = Alignment(vertical="center", wrap_text=(ci in (8,9)))
            cell.border = make_border()
            if ci == 6:
                cell.fill = PatternFill("solid", fgColor=sc)
                cell.font = Font(name="Calibri", bold=True, size=10, color=TEXT_WHITE)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.fill = PatternFill("solid", fgColor=bg)
        ws.row_dimensions[rn].height = 18

    ws.freeze_panes = "A3"
    ws.sheet_properties.tabColor = HEADER_BG
    buf = BytesIO(); wb.save(buf); buf.seek(0)
    fname = f"jobtracker_{user['username']}_{datetime.utcnow().strftime('%Y%m%d')}.xlsx"
    return send_file(buf, as_attachment=True, download_name=fname,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

# ---------------------------------------------------------------------------
# Real job search — Adzuna API proxy
# ---------------------------------------------------------------------------

@app.route("/api/search", methods=["GET"])
@require_auth
def search_jobs(user):
    app_id  = os.environ.get("ADZUNA_APP_ID", "")
    app_key = os.environ.get("ADZUNA_APP_KEY", "")
    if not app_id or not app_key:
        return jsonify({"error": "ADZUNA_APP_ID and ADZUNA_APP_KEY not configured on server"}), 503

    what     = request.args.get("q", "data scientist junior")
    where    = request.args.get("location", "France")
    country  = request.args.get("country", "fr")
    page     = request.args.get("page", "1")
    per_page = request.args.get("per_page", "10")

    params = urllib.parse.urlencode({
        "app_id":          app_id,
        "app_key":         app_key,
        "what":            what,
        "where":           where,
        "results_per_page": per_page,
        "content-type":    "application/json",
        "sort_by":         "date",
    })
    url = f"https://api.adzuna.com/v1/api/jobs/{country}/search/{page}?{params}"

    try:
        with urllib.request.urlopen(url, timeout=10) as resp:
            raw = json.loads(resp.read().decode())
    except Exception as e:
        return jsonify({"error": f"Adzuna request failed: {str(e)}"}), 502

    jobs = []
    for r in raw.get("results", []):
        jobs.append({
            "id":          r.get("id", ""),
            "title":       r.get("title", ""),
            "company":     r.get("company", {}).get("display_name", ""),
            "location":    r.get("location", {}).get("display_name", ""),
            "description": r.get("description", ""),
            "url":         r.get("redirect_url", ""),
            "salary_min":  r.get("salary_min"),
            "salary_max":  r.get("salary_max"),
            "created":     r.get("created", "")[:10],
            "category":    r.get("category", {}).get("label", ""),
        })

    return jsonify({
        "jobs":  jobs,
        "total": raw.get("count", 0),
        "page":  int(page),
    })


# ---------------------------------------------------------------------------
# Health check
# ---------------------------------------------------------------------------

@app.route("/api/health", methods=["GET"])
def health():
    try:
        result = query("SELECT COUNT(*) as n FROM users", fetchone=True)
        count = result["n"] if result else 0
        return jsonify({
            "status": "ok",
            "db": "postgresql" if USE_PG else "sqlite",
            "database_url_set": bool(DATABASE_URL),
            "users": count,
        })
    except Exception as e:
        return jsonify({"status": "error", "error": str(e)}), 500

# ---------------------------------------------------------------------------
# Entrypoint
# ---------------------------------------------------------------------------

try:
    init_db()
    print(f"✅ DB initialized ({'PostgreSQL' if USE_PG else 'SQLite'})")
except Exception as e:
    print(f"❌ init_db() FAILED: {e}")
    raise

if __name__ == "__main__":
    app.run(debug=True, port=int(os.environ.get("PORT", 5000)))
